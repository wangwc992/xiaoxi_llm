from pptx import Presentation


def extract_text_and_tables_from_pptx(file_path):
    prs = Presentation(file_path)
    extracted_content = []

    for slide in prs.slides:
        slide_content = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                # 提取文本内容
                slide_content.append(shape.text.strip())

            if shape.has_table:
                # 提取表格内容
                table = shape.table
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    slide_content.append("\t".join(row_text).strip())  # 用制表符分隔单元格内容
        if slide_content:
            extracted_content.append(" ".join(slide_content))

    return "\n\n".join(extracted_content)


# 使用示例
pptx_file = r"C:\Users\wishfyc\Desktop\2023032019105405803013.pptx"  # 确保路径正确
extracted_content = extract_text_and_tables_from_pptx(pptx_file)
print(extracted_content)
