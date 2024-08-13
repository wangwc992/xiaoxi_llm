import openpyxl
import fitz
from docx import Document
from pptx import Presentation

from app.common.core.paddle_ocr import process_pdf
from app.common.utils.file_utils import download_file, delete_file


def get_file_extension(file_obj):
    # 尝试从文件名中获取扩展名
    try:
        return file_obj.split('.')[-1].lower()
    except AttributeError:
        return ""


def docToString(file_obj):
    text = ""
    try:
        doc = Document(file_obj)
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
    except Exception as e:
        print("Error:", e)
    return text.strip()


def xlsToString(file_obj):
    result_text = ""
    try:
        wb = openpyxl.load_workbook(file_obj, data_only=True)
        for sheet_name in wb.sheetnames:
            sheet = wb[sheet_name]
            result_text += f"Sheet: {sheet_name}"
            merged_cells = sheet.merged_cells.ranges  # Get merged cell ranges
            for row in sheet.iter_rows():
                text = ""
                for cell in row:
                    if cell is not None and cell.value is not None:
                        text = str(cell.value).replace("None", "").strip()
                        if text:
                            result_text += text + " "
                    else:
                        # Check if the cell is part of a merged cell range
                        for merged_range in merged_cells:
                            if cell is not None and sheet.cell(row=cell.row,
                                                               column=cell.column).coordinate in merged_range:
                                # Get the value of the top-left cell of the merged range
                                top_left_cell = sheet.cell(row=merged_range.min_row, column=merged_range.min_col)
                                text = str(top_left_cell.value).replace("None", "").strip()
                                if text:
                                    result_text += text + " "
                                break
                if text:
                    result_text += "\n"
    except Exception as e:
        print("Error:", e)
    return result_text.strip()


def pptToString(file_path):
    prs = Presentation(file_path)
    extracted_content = []

    for slide in prs.slides:
        slide_content = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                # 提取文本内容
                slide_content.append(shape.text.strip())

            if shape.has_table:
                # 提取表格内容
                table = shape.table
                for row in table.rows:
                    row_text = [cell.text for cell in row.cells]
                    slide_content.append("\t".join(row_text).strip())  # 用制表符分隔单元格内容
        if slide_content:
            extracted_content.append(" ".join(slide_content))

    return "\n\n".join(extracted_content)


def pdfToString(file_path):
    doc = fitz.open(file_path)
    text = ""
    for page_num in range(min(len(doc), 3)):
        page = doc.load_page(page_num)
        try:
            page_text = page.get_text()
        except AssertionError:
            page_text = ""
            # 可以在此处添加其他错误处理逻辑
        if page_text.strip():
            text += page_text
        else:
            text = process_pdf(file_path)
    return text.strip()


def fileToString(file_path, file_type=None):
    if file_type == "pdf":
        text = pdfToString(file_path)
    elif file_type == "jpg" or file_type == "png" or file_type == "gif" or file_type == "jpeg":
        text = process_pdf(file_path)
    elif file_type == "doc" or file_type == "docx":
        text = docToString(file_path)
    elif file_type == "xls" or file_type == "xlsx":
        text = xlsToString(file_path)
    elif file_type in "pptx":
        text = pptToString(file_path)
    else:
        text = ""
    return text.strip()


def urlToText(url):
    file_path = download_file(url, "downloads")
    fileType = get_file_extension(file_path)
    text = fileToString(file_path, fileType)
    # 删除临时文件
    delete_file(file_path)
    return text
